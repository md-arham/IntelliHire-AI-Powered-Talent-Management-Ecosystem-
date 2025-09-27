import re
import uuid
from datetime import datetime
from io import BytesIO
from typing import Dict, Any
import pdfplumber
import pymupdf
from pydantic import ValidationError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select
import os
from app.database.models import (
    Certification,
    CompanyExperience,
    Education,
    Language,
    PersonalInfo,
    Projects,
    Skill,
    SocialProfile,
    WorkExperience,
)
from app.schemas.resume_schemas import ResumeData
from app.utils.logger_config import logger
from app.utils.settings import async_huggingface_client
from docx import Document


def parse_date_safe(date_str):
    if date_str:
        try:
            return datetime.strptime(date_str, "%Y-%m-%d").date()
        except ValueError:
            return None
    return None

def parse_dates(dates_str):
    try:
        start, end = dates_str.split("-")
        start_date = datetime.strptime(start.strip(), "%Y").date()
        end = end.strip().lower()
        end_date = (
            None
            if end in ["present", "ongoing"]
            else datetime.strptime(end, "%Y").date()
        )
        return start_date, end_date
    except Exception:
        return None, None


def extract_grade(details_list):
    for item in details_list:
        if any(keyword in item.lower() for keyword in ["cgpa", "gpa", "percentage"]):
            return item.split(":")[-1].strip()
    return None

async def store_personal_info_V2(data, db: AsyncSession):
    try:
        info = PersonalInfo(
            id=uuid.uuid4(),
            version_id=data.get("version_id"),
            student_id=data.get("student_id"),
            full_name=data.get("name", ""),
            email=data.get("email", ""),
            phone=data.get("phone", ""),
            location="",
            bio="",
        )
        db.add(info)
        await db.commit()
        await db.refresh(info)
        logger.info("Created Personal Info IN DB")
        return info.id
    except Exception as e:
        logger.info(f"Error: {str(e)}")
        return None


async def store_company_and_work_experience_V2(
    json_data, personal_info_id, db: AsyncSession
):
    result = await db.execute(select(PersonalInfo).filter_by(id=personal_info_id))
    personal_info = result.scalar_one()
    experiences = json_data.get("experience", [])

    for exp in experiences:
        company = exp.get("company")
        role = exp.get("role")
        dates = exp.get("dates")

        stmt = select(CompanyExperience).filter_by(
            user=personal_info,
            company_name=company,
            your_position=role,
            dates=dates,
        )
        result = await db.execute(stmt)
        existing_company_exp = result.scalar_one_or_none()

        if existing_company_exp:
            existing_company_exp.company_description = (
                " ".join(exp.get("details", []))
                or existing_company_exp.company_description
            )

            stmt = select(WorkExperience).filter_by(
                company=existing_company_exp, position_title=role
            )
            result = await db.execute(stmt)
            existing_work_exp = result.scalar_one_or_none()

            if existing_work_exp:
                existing_work_exp.description = (
                    " ".join(exp.get("details", [])) or existing_work_exp.description
                )
            else:
                new_work_exp = WorkExperience(
                    id=uuid.uuid4(),
                    company=existing_company_exp,
                    position_title=role,
                    company_name=company,
                    duration=dates,
                    description=" ".join(exp.get("details", [])),
                )
                db.add(new_work_exp)
        else:
            company_exp = CompanyExperience(
                id=uuid.uuid4(),
                user=personal_info,
                company_name=company,
                your_position=role,
                dates=dates,
                company_website=None,
                company_location=None,
                company_description=" ".join(exp.get("details", [])),
                company_logo=None,
            )
            db.add(company_exp)
            await db.flush()

            work_exp = WorkExperience(
                id=uuid.uuid4(),
                company=company_exp,
                position_title=role,
                company_name=company,
                duration=dates,
                description=" ".join(exp.get("details", [])),
            )
            db.add(work_exp)
    await db.commit()
    logger.info("Added Company and Work Experience")


async def store_skills_from_json_V2(json_data, personal_info_id, db: AsyncSession):
    result = await db.execute(select(PersonalInfo).filter_by(id=personal_info_id))
    personal_info = result.scalar_one()
    skills_data = json_data.get("skills", {})

    def prettify_category(cat_key):
        return cat_key.replace("_", " ").title()

    for category_key, skill_list in skills_data.items():
        category = prettify_category(category_key)
        for skill in skill_list:
            skill = skill.strip()
            if not skill:
                continue
            stmt = select(Skill).filter_by(user=personal_info, skill_name=skill)
            result = await db.execute(stmt)
            existing = result.scalar_one_or_none()
            if not existing:
                db.add(
                    Skill(
                        id=uuid.uuid4(),
                        user=personal_info,
                        skill_name=skill,
                        category=category,
                    )
                )

    await db.commit()
    logger.info("Added Store Skills")


async def store_certifications_from_json_V2(
    json_data, personal_info_id, db: AsyncSession
):
    result = await db.execute(select(PersonalInfo).filter_by(id=personal_info_id))
    personal_info = result.scalar_one()
    certs = json_data.get("extra_fields", {}).get("certifications", [])

    for cert in certs:
        name = cert.get("certification_name")
        issued_by = cert.get("issued_by")

        stmt = select(Certification).filter_by(
            user=personal_info, certification_name=name, issued_by=issued_by
        )
        result = await db.execute(stmt)
        existing = result.scalar_one_or_none()

        if existing:
            existing.issue_date = (
                parse_date_safe(cert.get("issue_date")) or existing.issue_date
            )
            existing.expiration_date = (
                parse_date_safe(cert.get("expiration_date")) or existing.expiration_date
            )
        else:
            db.add(
                Certification(
                    id=uuid.uuid4(),
                    user=personal_info,
                    certification_name=name,
                    issued_by=issued_by,
                    issue_date=parse_date_safe(cert.get("issue_date")),
                    expiration_date=parse_date_safe(cert.get("expiration_date")),
                )
            )

    await db.commit()
    logger.info("Added Certifications into DB")


async def store_social_profiles_from_json_V2(
    json_data: dict, personal_info_id, db: AsyncSession
):
    result = await db.execute(select(PersonalInfo).filter_by(id=personal_info_id))
    personal_info = result.scalar_one()
    stmt = select(SocialProfile).filter_by(user=personal_info)
    result = await db.execute(stmt)
    existing = result.scalar_one_or_none()

    if existing:
        existing.personal_website = json_data.get(
            "personal_website", existing.personal_website
        )
        existing.linkedin = json_data.get("linkedin", existing.linkedin)
        existing.github = json_data.get("github", existing.github)
        existing.twitter = json_data.get("twitter", existing.twitter)
    else:
        db.add(
            SocialProfile(
                id=uuid.uuid4(),
                user=personal_info,
                personal_website=json_data.get("personal_website"),
                linkedin=json_data.get("linkedin"),
                github=json_data.get("github"),
                twitter=json_data.get("twitter"),
            )
        )

    await db.commit()
    logger.info("Added social profiles into db")


async def store_education_from_json_V2(json_data, personal_info_id, db: AsyncSession):
    result = await db.execute(select(PersonalInfo).filter_by(id=personal_info_id))
    personal_info = result.scalar_one()
    education_data = json_data.get("education", [])

    for edu in education_data:
        institution = edu.get("institution")
        degree = edu.get("degree")

        stmt = select(Education).filter_by(
            user=personal_info, institution_name=institution, degree=degree
        )
        result = await db.execute(stmt)
        existing = result.scalar_one_or_none()

        start_date, end_date = parse_dates(edu.get("dates", ""))
        grade = extract_grade(edu.get("details", []))
        description = " ".join(edu.get("details", []))

        if existing:
            existing.start_date = start_date or existing.start_date
            existing.end_date = end_date or existing.end_date
            existing.grade = grade or existing.grade
            existing.description = description or existing.description
        else:
            db.add(
                Education(
                    id=uuid.uuid4(),
                    user=personal_info,
                    institution_name=institution,
                    degree=degree,
                    field_of_study="",
                    start_date=start_date,
                    end_date=end_date,
                    grade=grade,
                    description=description,
                )
            )

    await db.commit()
    logger.info("Added Education into DB")


async def store_languages_from_json_V2(json_data, personal_info_id, db: AsyncSession):
    result = await db.execute(select(PersonalInfo).filter_by(id=personal_info_id))
    personal_info = result.scalar_one()
    language_list = json_data.get("extra_fields", {}).get("languages", [])

    for lang in language_list:
        stmt = select(Language).filter_by(user=personal_info, language_name=lang)
        result = await db.execute(stmt)
        existing = result.scalar_one_or_none()

        if not existing:
            db.add(
                Language(
                    id=uuid.uuid4(),
                    user=personal_info,
                    language_name=lang,
                    proficiency_level="Fluent",
                )
            )

    await db.commit()
    logger.info("Added Languages into DB")


async def store_projects_from_json_V2(json_data, personal_info_id, db: AsyncSession):
    result = await db.execute(select(PersonalInfo).filter_by(id=personal_info_id))
    personal_info = result.scalar_one()
    projects = json_data.get("projects", [])

    for proj in projects:
        title = proj.get("title")
        stmt = select(Projects).filter_by(user=personal_info, title=title)
        result = await db.execute(stmt)
        existing = result.scalar_one_or_none()

        if existing:
            existing.dates = proj.get("dates", existing.dates)
            existing.details = " ".join(proj.get("details", [])) or existing.details
        else:
            db.add(
                Projects(
                    id=uuid.uuid4(),
                    user=personal_info,
                    title=title,
                    dates=proj.get("dates"),
                    details=" ".join(proj.get("details", [])),
                )
            )

    await db.commit()
    logger.info("Added Projects")


async def extract_resume_data_v2(file_stream: BytesIO, filename: str) -> Dict[str, Any]:
    """
    Extract resume data from PDF or DOCX files and return structured information.

    Args:
        file_stream: BytesIO stream of the file
        filename: Name of the file including extension

    Returns:
        Dict containing structured resume data with layout analysis
    """
    try:
        extension = os.path.splitext(filename)[-1].lower()

        if extension == ".pdf":
            return await _process_pdf_resume(file_stream)
        elif extension == ".docx":
            return await _process_docx_resume(file_stream)
        else:
            raise ValueError(f"Unsupported file format: {extension}")

    except Exception as e:
        raise RuntimeError(f"Failed to extract resume data: {str(e)}")


async def _process_pdf_resume(file_stream: BytesIO) -> Dict[str, Any]:
    """Process PDF resume and extract text with layout analysis."""
    # Extract text and detect basic features with pdfplumber
    full_text, has_images, has_tables = _extract_pdf_text_and_features(file_stream)

    # Get structured data from LLM
    structured_data = await descriptions_v2(resume_data=full_text)

    # Analyze advanced PDF features with pymupdf
    file_stream.seek(0)
    doc = pymupdf.open(stream=file_stream.read(), filetype="pdf")

    layout_features = _analyze_pdf_layout(doc)

    # Combine all data
    structured_data.update(
        {"image_present": has_images, "tables_present": has_tables, **layout_features}
    )

    doc.close()
    return structured_data


async def _process_docx_resume(file_stream: BytesIO) -> Dict[str, Any]:
    """Process DOCX resume and extract text with feature analysis."""
    file_stream.seek(0)
    document = Document(file_stream)

    # Extract text
    full_text = "\n".join(para.text for para in document.paragraphs)
    logger.info(f"Full text: {full_text}")

    # Get structured data from LLM
    structured_data = await descriptions_v2(resume_data=full_text)

    # Analyze document features
    features = _analyze_docx_features(document)
    structured_data.update(features)

    return structured_data


def _extract_pdf_text_and_features(file_stream: BytesIO) -> tuple[str, bool, bool]:
    """Extract text and detect images/tables from PDF using pdfplumber."""
    full_text = ""
    has_images = False
    has_tables = False

    with pdfplumber.open(file_stream) as pdf:
        for page in pdf.pages:
            # Extract text
            page_text = page.extract_text()
            if page_text:
                full_text += page_text + "\n"

            # Check for meaningful images (size threshold to avoid tiny decorative elements)
            if page.images:
                has_images = any(
                    img["width"] > 50 and img["height"] > 50 for img in page.images
                )

            # Check for tables
            if page.extract_tables():
                has_tables = True

    logger.info(f"Full text: {full_text}")
    return full_text, has_images, has_tables


def _analyze_pdf_layout(doc) -> Dict[str, bool]:
    """Analyze PDF layout features using pymupdf."""
    has_vectors = False
    has_icon_fonts = False

    # Check each page for vectors and icon fonts
    for page in doc:
        # Check for vector graphics
        if page.get_drawings():
            has_vectors = True

        # Check for icon fonts
        blocks = page.get_text("dict")["blocks"]
        for block in blocks:
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    font_name = span["font"]
                    text = span["text"]

                    # Detect icon fonts by name or Unicode range
                    if "FontAwesome" in font_name or (
                        text and any(ord(char) > 10000 for char in text)
                    ):
                        has_icon_fonts = True

    # Detect multi-column layout
    has_multi_column = _detect_multi_column_layout(doc)

    return {
        "icons_present": has_icon_fonts or has_vectors,
        "multi_column": has_multi_column,
    }


def _detect_multi_column_layout(doc, min_column_gap: int = 100) -> bool:
    """Detect if PDF has multi-column layout by analyzing text block positions."""
    for page_num in range(len(doc)):
        page = doc[page_num]
        blocks = page.get_text(
            "blocks"
        )  # (x0, y0, x1, y1, "text", block_no, block_type)

        if len(blocks) < 2:
            continue

        # Sort blocks by x-coordinate and get unique x positions
        x_coords = sorted(set(block[0] for block in blocks))

        # Check for significant gaps between columns
        for i in range(1, len(x_coords)):
            if x_coords[i] - x_coords[i - 1] > min_column_gap:
                logger.info(f"Page {page_num + 1}: Detected multi-column layout.")
                return True

    logger.info("No multi-column layout detected.")
    return False


def _analyze_docx_features(document) -> Dict[str, bool]:
    """Analyze DOCX document features."""
    # Check for images
    has_images = any("image" in rel.target_ref for rel in document.part.rels.values())

    # Check for tables
    has_tables = len(document.tables) > 0

    # Check for icon characters (high Unicode values)
    has_icons = any(
        any(ord(char) > 10000 for char in para.text)
        for para in document.paragraphs
        if para.text
    )

    # Check for multi-column layout
    has_multi_column = any(
        hasattr(section, "columns") and getattr(section.columns, "count", 1) > 1
        for section in document.sections
    )

    return {
        "image_present": has_images,
        "tables_present": has_tables,
        "icons_present": has_icons,
        "multi_column": has_multi_column,
    }


async def descriptions_v2(resume_data: str) -> Dict[str, Any]:
    """
    Extract structured resume data using LLM.

    Args:
        resume_data: Raw text extracted from resume

    Returns:
        Dictionary containing structured resume information
    """
    sys_prompt = _get_system_prompt()

    messages = [
        {"role": "system", "content": sys_prompt},
        {"role": "user", "content": f"Resume text: ` {resume_data}`"},
    ]

    # Get completion from LLM
    completion = await async_huggingface_client.chat.completions.create(
        model="tgi",
        messages=messages,
        top_p=0.1,
        temperature=0,
        max_tokens=2000,
        stream=False,
    )

    content = completion.choices[0].message.content
    logger.info(f"***************Async Content***********: {content}")

    return _parse_llm_response(content)


def _get_system_prompt() -> str:
    """Return the system prompt for resume parsing."""
    return """
You are an AI specialized in resume parsing and ATS scoring.
Given the following resume text, extract and structure the information into this JSON format without hallucinating:

json:{
    "name": "",
    "phone": "",
    "email": "",
    "linkedin": "",
    "github": "",
    "twitter": "",
    "personal_website": "",
    "objective": "",
    "education": [
        {
            "degree": "",
            "institution": "",
            "dates": "",
            "details": [""]
        }
    ],
    "experience": [
        {
            "role": "",
            "company": "",
            "dates": "",
            "details": [""]
        }
    ],
    "projects": [
        {
            "title": "",
            "dates": "",
            "details": [""]
        }
    ],
    "project_clarity": true,
    "skills": {
        //Only process content explicitly under the Skills section.
        //If the Skills section uses a key-value structure (e.g., Programming Languages: Python, Java), then:
            //Use the exact same keys from the resume as the skill category names.
            //Use the listed skills as a flat list of strings under each corresponding key.
            //Do NOT rename keys (e.g., don't replace "Programming Languages" with "Core Skills").
        //If the skills section is not written as key-value pairs (i.e., a simple list or a space-separated or comma-separated group of skills under a single heading like "Professional Skills" or "Skills"), then:

            //Do not treat any individual skill as a key or header.

            //Instead, classify all skills into exactly two categories:

                //"Core Skills": Include all functional/technical, business, or domain-specific competencies .

                //"Soft Skills": Include all interpersonal, communication, leadership, and time-management-related skills.

                //Return the result as a flat JSON dictionary with only two keys: "Core Skills" and "Soft Skills".
        //All skill category names must be in Title Case With Spaces.
        //All skill lists must be flat arrays of strings. Do not nest any dictionaries or categories inside.
    },
    "achievements": [""],
    "publications": [""],
    "extra_fields": {
        "certifications": [
            {
                "certification_name": "",
                "issued_by": "",
                "issue_date": "",
                "expiration_date": ""
            }
        ],
        "languages": [""],
        "courses": [""]
    },
    "valid_descriptions": true,
    "resume_improvements": ""
}

Instructions:
    -Leave any missing or unavailable fields as empty strings "" or empty lists [].
    -project_clarity should be true only if all listed project descriptions are specific and complete. Set to false if they are vague or missing.
    -valid_descriptions should be false if any part of the resume contains unprofessional, placeholder, or meaningless text (e.g., "lorem ipsum", slang, or irrelevant filler).
    -For "resume_improvements":
        -Provide feedback after reviewing the professionalism and clarity of descriptions under each section.
        -Suggest clear, actionable improvements in string format with examples to improve the ATS score.
    - languages has to be a list of strings
    - courses has to be a list of strings

Return the completed JSON.
"""


def _parse_llm_response(content: str) -> Dict[str, Any]:
    """Parse and validate LLM response to extract JSON data."""
    try:
        # Try to find JSON block with ```json wrapper first
        json_match = re.search(r"```json\s*(\{.*?\})\s*```", content, re.DOTALL)

        # Fallback: Look for any JSON-like object
        if not json_match:
            json_match = re.search(r"(\{.*\})", content, re.DOTALL)

        if not json_match:
            raise ValueError("No valid JSON object found in the response.")

        json_str = json_match.group(1)

        # Validate and parse using your ResumeData model
        structured_description = ResumeData.model_validate_json(json_str)
        return structured_description.model_dump()

    except (ValidationError, ValueError) as e:
        logger.error(f"Failed to extract or store resume data from LLM: {str(e)}")
        raise RuntimeError(f"Failed to extract resume data: {str(e)}")
